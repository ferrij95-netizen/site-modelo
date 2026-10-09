"""Escolhe os 100 dentistas (movimento recente + site com 4+ problemas ou sem site), sem repetir a lista de clínicas.
Uso: python final.py selecao.json [clinicas.xlsx]"""
import json, re, math, sys
from collections import Counter, defaultdict
R = json.load(open('saida/resultado.json'))
UN = {'minuto': 0, 'hora': 0, 'dia': 1, 'semana': 7, 'mês': 30, 'mes': 30, 'meses': 30, 'ano': 365}
def dias(s):
    m = re.search(r'(um|uma|\d+)\s+(minuto|hora|dia|semana|m[eê]s|meses|ano)', s.lower())
    if not m: return None
    return (1 if m.group(1) in ('um', 'uma') else int(m.group(1))) * UN[m.group(2)]
def campo(ficha, rx):
    m = re.search(rx, ficha or '')
    return m.group(0).strip() if m else ''
CAT_OK = re.compile(r'dentist|odontol|ortodon|implant|endodon|periodon', re.I)
NOME_OK = re.compile(r'odonto|dentist|dental|sorriso|ortodon|implant|\bdra?\.? ', re.I)
FORA = re.compile(r'laborat|loja|suprimento|distribuid|equipamento|veterin|faculdade|escola|curso|hospital|radiolog|documenta[cç][aã]o odonto', re.I)
def e_dent(r):
    if FORA.search((r.get('nome') or '') + ' ' + (r.get('categoria') or '')): return False
    return bool(CAT_OK.search(r.get('categoria') or ''))
R = [r for r in R if e_dent(r)]
# não repetir quem já está na lista de clínicas (abas Clínicas e Reservas)
dig = lambda t: re.sub(r'\D', '', t or '')[-8:]
nrm = lambda t: re.sub(r'[^a-z0-9]', '', (t or '').lower())
dom = lambda u: re.sub(r'^https?://(www\.)?', '', (u or '').lower()).split('/')[0]
ja_tel, ja_nome, ja_dom = set(), set(), set()
if len(sys.argv) > 2:
    from openpyxl import load_workbook
    for ws in load_workbook(sys.argv[2], read_only=True).worksheets:
        for row in list(ws.iter_rows(values_only=True))[1:]:
            ja_nome.add(nrm(row[1]))
            if row[12]: ja_tel.add(dig(str(row[12])))
            if row[9] and str(row[9]).startswith('http') and 'instagram' not in str(row[9]): ja_dom.add(dom(row[9]))
antes = len(R)
R = [r for r in R if nrm(r['nome']) not in ja_nome and not (r.get('telefone') and dig(r['telefone']) in ja_tel)
     and not (r.get('tipo_site') == 'site próprio' and dom(r['site']) in ja_dom)]
print('fora da lista de clínicas:', len(R), 'de', antes)
out = []
for r in R:
    d = r.get('detalhe', {})
    ds = [x for x in (dias(s) for s in d.get('datas') or []) if x is not None]
    r['ult_dias'] = min(ds) if ds else None
    r['ult_txt'] = next((s for s in d.get('datas') or [] if dias(s) == r['ult_dias']), '')
    cert = 'ERR_CERT' in (d.get('erro_site') or '')
    P = [p for p in r.get('problemas', []) if not (p.startswith('lento (99s') or p == 'certificado/HTTPS com erro')]
    if cert: P.insert(0, 'navegador mostra aviso de "site não seguro" (certificado inválido)')
    elif r.get('erro') and not r.get('site_ok'): P.insert(0, 'site fora do ar ou não abre')
    if d.get('erros_js'): P.append(f'{len(d["erros_js"])} erro(s) de programação ao abrir a página')
    if (d.get('n_falhas') or 0) >= 3: P.append(f'{d["n_falhas"]} arquivos ou imagens que não carregam')
    if d.get('larg') and d.get('tela') and d['larg'] > d['tela'] + 20 and 'não adaptado para celular' not in P: P.append('página mais larga que a tela do celular')
    if (d.get('carga_s') or 0) > 8: P.append(f'demora {d["carga_s"]:.0f}s para abrir no celular')
    if (d.get('status_cel') or 0) >= 500: P.append(f'erro {d["status_cel"]} do servidor')
    r.setdefault('nota_g', r.get('nota'))
    r['sem_site'] = r.get('tipo_site') != 'site próprio'
    if r['sem_site']:
        P = ['não tem site próprio' + (f' (só {dom(r["site"])})' if r['site'] else ' (só a ficha do Google)')]
    r['P'] = P
    r['nota'] = 10 if r['sem_site'] else min(10, len(P))
    r['abre'] = bool(r.get('site_ok')) or cert
    r['endereco'] = campo(d.get('ficha'), r'[^\n]*\d{5}-\d{3}[^\n]*')
    out.append(r)
def score(r):
    rec = 3 if r['ult_dias'] <= 30 else 2 if r['ult_dias'] <= 90 else 1
    # site ruim vem antes de sem site; dentro de cada grupo, mais problemas, mais recente, mais avaliações
    return (0 if r['sem_site'] else 30) + (0 if r['sem_site'] else r['nota'] * 2) + rec + math.log10(1 + r['avaliacoes'])
ativo = lambda r: r['ult_dias'] is not None and r['ult_dias'] <= 180
ok = [r for r in out if ativo(r) and ((not r['sem_site'] and r['abre'] and r['nota'] >= 4) or r['sem_site'])]
ok.sort(key=score, reverse=True)
print('elegíveis', len(ok), 'sem site', sum(r['sem_site'] for r in ok), Counter(r['nota'] for r in ok if not r['sem_site']))
por = Counter(); top, res = [], []
for r in ok:
    if len(top) < 100 and por[r['cidade']] < 10:
        top.append(r); por[r['cidade']] += 1
    else:
        res.append(r)
# reservas: elegíveis que sobraram + ativas com nota 3 ou site fora do ar
extra = [r for r in out if r not in ok and ativo(r) and not r['sem_site'] and r['nota'] >= 3]
extra.sort(key=score, reverse=True)
res = (res + extra)[:100]
print('top', len(top), 'sem site', sum(r['sem_site'] for r in top), por, 'reservas', len(res))
json.dump({'top': top, 'res': res}, open(sys.argv[1], 'w'), ensure_ascii=False)
