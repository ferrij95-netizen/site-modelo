"""Planilha + documentos do canvas (lista construtoras) a partir de final.py. Uso: python planilha.py selecao.json saida.xlsx pasta_docs"""
import json, sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
S = json.load(open(sys.argv[1])); xlsx = sys.argv[2]; docsdir = sys.argv[3]
PORTE = {'01': 'Microempresa', '03': 'Pequeno porte (EPP)', '05': 'Médio/grande', '00': ''}
COLS = [('#', 5), ('Construtora', 34), ('Cidade', 20), ('Pontuação bom pagador (0-100)', 12), ('Por que a pontuação', 70),
        ('Situação do site', 14), ('Site', 32), ('Nota de site ruim (0-10)', 10), ('Problemas do site', 60),
        ('Fotos no Google', 9), ('Fotos do proprietário', 9), ('Obras no site', 40), ('Avaliações', 10), ('Nota Google', 8),
        ('Última avaliação', 14), ('Anos de CNPJ', 8), ('Capital social (R$)', 14), ('Porte', 16), ('CNPJ', 18), ('Razão social', 36),
        ('Como o CNPJ foi achado', 12), ('Telefone', 17), ('WhatsApp', 15), ('E-mail', 28), ('Endereço', 45), ('Google Maps', 14)]
def tel(r): return r.get('telefone') or r.get('tel_site') or ''
def wa(r): return ('+' + r['whatsapp']) if r.get('whatsapp') else ''
def email(r): return r.get('email') or (r['rf'] or {}).get('email', '')
def fcnpj(c): return f'{c[:2]}.{c[2:5]}.{c[5:8]}/{c[8:12]}-{c[12:]}' if c else ''
def end(r):
    e = r['rf']
    return f"{e['end']}, {e['bairro']}, {e['cidade'].title()}-{e['uf']}" if e else ''
def linha(i, r):
    e = r['rf'] or {}
    return [i, r['nome'], r['cidade'], r['score'], '; '.join(r['S']), r['situacao_site'], r['site'] if r['tipo_site'] == 'site próprio' else (r['site'] or ''),
            r['nota_site'], '; '.join(r['P']), r['fotos'], r['fotos_dono'], '; '.join(r.get('obras_site') or []), r['avaliacoes'], r.get('nota'),
            r['ult_txt'], e.get('anos', ''), round(e['capital']) if e.get('capital') else '', PORTE.get(e.get('porte'), ''), fcnpj(e.get('cnpj', '')),
            e.get('razao', ''), r['rf_como'], tel(r), wa(r), email(r), end(r), r['maps']]
wb = Workbook()
for nome, lst in (('500 construtoras', S['top']), ('Reservas', S['res'])):
    ws = wb.active if nome.startswith('500') else wb.create_sheet(nome)
    ws.title = nome
    ws.append([c for c, _ in COLS])
    for i, r in enumerate(lst, 1):
        ws.append(linha(i, r))
        if r['site']: ws.cell(i + 1, 7).hyperlink = r['site']
        ws.cell(i + 1, 26).hyperlink = r['maps']; ws.cell(i + 1, 26).value = 'abrir no Maps'
        ws.cell(i + 1, 17).number_format = '#,##0'
    for j, (_, w) in enumerate(COLS, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
        c = ws.cell(1, j); c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='182644')
        c.alignment = Alignment(wrap_text=True, vertical='center')
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'C2'; ws.auto_filter.ref = ws.dimensions
ws = wb.create_sheet('Como foi feito')
for t in json.load(open(sys.argv[4])) if len(sys.argv) > 4 else []:
    ws.append([t])
ws.column_dimensions['A'].width = 150
for c in ws['A']: c.alignment = Alignment(wrap_text=True)
ws['A1'].font = Font(bold=True, size=14)
wb.save(xlsx)
os.makedirs(docsdir, exist_ok=True)
docs = []
for i, r in enumerate(S['top'], 1):
    e = r['rf'] or {}
    docs.append({'id': f'constr-{i:03d}', 'data': {
        'lista': 'construtoras', 'nome': r['nome'], 'site': r['site'] if r['tipo_site'] == 'site próprio' else '', 'cidade': r['cidade'],
        'segmento': r.get('categoria') or 'Construtora',
        'notas': f"Bom pagador {r['score']}/100: " + '; '.join(r['S']) + '. Site: ' + '; '.join(r['P']),
        'status': 'novo', 'rank': i, 'triagem': '',
        'contato': {'telefone': tel(r), 'email': email(r), 'whatsapp': wa(r)},
        'extra': {'notaGoogle': r.get('nota'), 'avaliacoes': r['avaliacoes'], 'notaSite': r['nota_site'], 'maps': r['maps'],
                  'endereco': end(r), 'situacaoSite': r['situacao_site'], 'problemas': r['P'], 'ultimaAvaliacao': r['ult_txt'],
                  'bomPagador': r['score'], 'motivosBomPagador': r['S'], 'fotosGoogle': r['fotos'], 'fotosDono': r['fotos_dono'],
                  'obrasSite': r.get('obras_site') or [], 'cnpj': fcnpj(e.get('cnpj', '')), 'razaoSocial': e.get('razao', ''),
                  'anosCnpj': e.get('anos'), 'capitalSocial': e.get('capital'), 'porte': PORTE.get(e.get('porte'), '')},
        'criadoEm': 1791520000000 - i}})
for k in range(0, len(docs), 50):
    json.dump([{'op': 'set', 'collection': 'empresas', 'doc_id': d['id'], 'data': d['data']} for d in docs[k:k + 50]],
              open(f'{docsdir}/lote-{k // 50:02d}.json', 'w'), ensure_ascii=False)
print('ok', len(docs))
