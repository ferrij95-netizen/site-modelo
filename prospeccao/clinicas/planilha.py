"""Monta a planilha de clínicas: primeiro as sem site, depois os sites com nota pior."""
import json, re, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

UF = {'São Paulo': 'SP', 'Rio de Janeiro': 'RJ', 'Belo Horizonte': 'MG', 'Brasília': 'DF', 'Curitiba': 'PR',
      'Porto Alegre': 'RS', 'Salvador': 'BA', 'Recife': 'PE', 'Fortaleza': 'CE', 'Goiânia': 'GO', 'Campinas': 'SP',
      'Florianópolis': 'SC', 'Ribeirão Preto': 'SP', 'Londrina': 'PR', 'Joinville': 'SC', 'Caxias do Sul': 'RS',
      'Uberlândia': 'MG', 'Sorocaba': 'SP', 'Natal': 'RN', 'Maceió': 'AL', 'Campo Grande': 'MS', 'Cuiabá': 'MT',
      'Vitória': 'ES', 'Juiz de Fora': 'MG', 'Maringá': 'PR', 'Santos': 'SP', 'São José dos Campos': 'SP',
      'Teresina': 'PI', 'João Pessoa': 'PB', 'Belém': 'PA'}
# hospitais, laboratórios e redes grandes não são o alvo
FORA = re.compile(r'hospital|laborat|unimed|hapvida|amil|sorridents|odontocompany|oral sin|orthopride|odonto excellence|'
                  r'dr\.? ?consulta|clínica sim|clinica sim|amo saúde|cartão de todos|espaçolaser|espaco laser|'
                  r'mais top estética|onodera|emagrecentro|pronto[- ]?socorro|upa |ubs |posto de saúde|sesi|senac|'
                  r'drogaria|farmácia|farmacia|petz|veterin|pet |shopping|centro comercial|amorsa[uú]de|sódent|sodent|'
                  r'doutor dermato|dr\.? exame|medclin|grupo |rede |franquia|academia|salão|salao|barbearia|ótica|otica', re.I)
ESPEC = [('Odontologia', r'odonto|dent|ortodon|implant|sorriso'),
         ('Estética', r'estétic|estetic|harmoniza|beleza|depila|botox|spa\b|micropigment|laser'),
         ('Dermatologia', r'dermato'), ('Cirurgia plástica', r'plástic|plastic'),
         ('Ginecologia', r'gineco|obstet|mulher'), ('Fisioterapia', r'fisio|pilates|reabilita'),
         ('Médica', r'médic|medic|clínic|clinic|consult|saúde|saude|doutor|dr\.|dra\.')]

def especialidade(p):
    t = (p.get('categoria', '') + ' ' + p['nome'] + ' ' + ' '.join(p['buscas'])).lower()
    for nome, rx in ESPEC:
        if re.search(rx, (p.get('categoria', '') + ' ' + p['nome']).lower()):
            return nome
    for nome, rx in ESPEC:
        if re.search(rx, t):
            return nome
    return 'Médica'

L = json.load(open(sys.argv[1]))
L = [p for p in L if not FORA.search(p['nome'] + ' ' + p.get('categoria', ''))]
# nomes que se repetem em 3+ fichas são redes (cada unidade aparece "sem site" mas a rede tem site)
from collections import Counter, defaultdict
chave = lambda p: ' '.join(re.sub(r'[^a-z0-9 ]', '', p['nome'].lower()).split()[:2])
rep = Counter(chave(p) for p in L)
L = [p for p in L if rep[chave(p)] < 3]
ok = lambda p: (p['nota'] or 0) >= 4.3 and 15 <= p['avaliacoes'] <= 1500 and p['telefone']

sem = [p for p in L if p['tipo_site'] in ('sem site', 'rede social/portal') and ok(p)]
# site ruim: 4+ problemas, ou site que responde com erro (404, 500, certificado); bloqueios de robô (403/429) ficam de fora
quebrado = lambda p: not p.get('site_ok') and (str(p.get('status', '')).startswith(('404', '410', '5')) or p.get('erro') == 'SSLError')
ruim = [p for p in L if p['tipo_site'] == 'site próprio' and ok(p) and (quebrado(p) or (p.get('site_ok') and p.get('pontos', 0) >= 4))]
sem.sort(key=lambda p: (p['tipo_site'] != 'sem site', -p['avaliacoes']))
ruim.sort(key=lambda p: (quebrado(p), -p.get('pontos', 0), -p['avaliacoes']))
print('sem site', len(sem), 'site ruim', len(ruim), 'total filtrado', len(L))

def espalha(lst, n):
    """Pega n itens alternando entre cidades, para não concentrar tudo numa capital."""
    por = defaultdict(list)
    for p in lst: por[p['cidade']].append(p)
    out = []
    while len(out) < n and any(por.values()):
        for c in list(por):
            if por[c] and len(out) < n: out.append(por[c].pop(0))
    return out, [p for v in por.values() for p in v]

N = int(sys.argv[3]) if len(sys.argv) > 3 else 300
n_ruim = min(len(ruim), N // 5)
a1, r1 = espalha(sem, N - n_ruim)
a2, r2 = espalha(ruim, n_ruim)
lista = a1 + a2
resto = sorted(r1, key=lambda p: -p['avaliacoes'])[:300] + r2[:100]

def linha(i, p):
    if p['tipo_site'] == 'sem site':
        site, sit, nota, prob = 'sem site', 'Sem site', '', 'Só tem o perfil no Google'
    elif p['tipo_site'] == 'rede social/portal':
        site, sit, nota, prob = p['site'], 'Só rede social/portal', '', 'Não tem site próprio, só Instagram/Doctoralia/link'
    elif not p.get('site_ok'):
        site, sit, nota, prob = p['site'], 'Site com erro', 10, 'Site não abre direito (' + str(p.get('status') or p.get('erro') or 'erro') + ')'
    else:
        site, sit, nota, prob = p['site'], 'Site desatualizado', p['pontos'], '; '.join(p['problemas'])
    zap = p.get('whatsapp') or ''
    if not zap and re.search(r'\)\s?9\d{4}', p['telefone']):
        zap = 'provável (celular)'
    return [i, p['nome'], especialidade(p), p.get('categoria', ''), p['cidade'], UF.get(p['cidade'], ''),
            p['nota'], p['avaliacoes'], sit, site, nota, prob, p['telefone'] or p.get('tel_site', ''), zap,
            p.get('email', ''), p['maps']]

CAB = ['#', 'Nome', 'Especialidade', 'Categoria no Google', 'Cidade', 'UF', 'Nota Google', 'Avaliações',
       'Situação', 'Site', 'Nota de desatualização (0-10)', 'Problemas', 'Telefone', 'WhatsApp', 'E-mail', 'Google Maps']
LARG = [5, 38, 16, 24, 18, 5, 8, 10, 20, 34, 12, 60, 17, 16, 26, 30]

wb = Workbook()
for aba, dados in (('Clínicas', lista), ('Reservas', resto)):
    ws = wb.active if aba == 'Clínicas' else wb.create_sheet(aba)
    ws.title = aba
    ws.append(CAB)
    for i, p in enumerate(dados, 1):
        ws.append(linha(i, p))
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F4E5F')
        c.alignment = Alignment(wrap_text=True, vertical='center')
    for i, w in enumerate(LARG, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r in ws.iter_rows(min_row=2):
        r[11].alignment = Alignment(wrap_text=True, vertical='top')
        if r[15].value:
            r[15].hyperlink = r[15].value; r[15].value = 'abrir no Maps'; r[15].font = Font(color='0563C1', underline='single')
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = ws.dimensions
wb.save(sys.argv[2])
print('planilha', len(lista), Counter(linha(0, p)[8] for p in lista), Counter(especialidade(p) for p in lista))
print('cidades', Counter(p['cidade'] for p in lista).most_common())
