"""Gera a planilha (uma aba por marca + reservas + como foi feito) e os lotes de documentos do canvas.
Uso: python planilha.py saida/selecao.json <arquivo.xlsx> <pasta-dos-lotes>"""
import json, sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
S = json.load(open(sys.argv[1])); xlsx = sys.argv[2]; docsdir = sys.argv[3]
CHAVE = {'BMW': 'bmw', 'Mercedes-Benz': 'mercedes', 'Audi': 'audi', 'Porsche': 'porsche', 'Land Rover/Jaguar': 'landrover', 'Volvo': 'volvo'}
COLS = [('#', 5), ('Oficina', 32), ('Marca', 14), ('Por que é dessa marca', 26), ('Cidade', 17), ('Bairro', 18), ('Situação do site', 22),
        ('Site', 30), ('Nota de site ruim (0-10)', 10), ('Problemas do site', 50), ('Sinais para vender o sistema', 55),
        ('O que clientes dizem (trecho)', 60), ('Última avaliação', 14), ('Avaliações', 10), ('Nota Google', 9),
        ('Telefone', 17), ('WhatsApp', 15), ('E-mail', 26), ('Endereço', 40), ('Google Maps', 14)]
wa = lambda r: ('+' + r['whatsapp']) if r.get('whatsapp') else ''
tel = lambda r: r.get('telefone') or r.get('tel_site') or ''
def linha(i, r):
    return [i, r['nome'], r['marca'], r['evidencia'], r['cidade'], ', '.join(r['bairros'] or []), r['situacao'], r['site'], r['nota'],
            '; '.join(r['P']), '; '.join(r['S']), ' | '.join(r['trechos']), r['ult_txt'], r['avaliacoes'], r['nota_google'],
            tel(r), wa(r), r.get('email') or '', r.get('endereco') or '', 'abrir no Maps']
wb = Workbook(); wb.remove(wb.active)
abas = [(f'{m.split("/")[0]} ({len(v)})', v) for m, v in S['marcas'].items()] + [('Reservas', S['res'])]
for nome, lst in abas:
    ws = wb.create_sheet(nome[:31])
    ws.append([c for c, _ in COLS])
    for i, r in enumerate(lst, 1):
        ws.append(linha(i, r))
        if r['site']: ws.cell(i + 1, 8).hyperlink = r['site']
        ws.cell(i + 1, 20).hyperlink = r['maps']
    for j, (_, w) in enumerate(COLS, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
        c = ws.cell(1, j); c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='182644')
        c.alignment = Alignment(wrap_text=True, vertical='center')
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'C2'; ws.auto_filter.ref = ws.dimensions
tot = sum(len(v) for v in S['marcas'].values())
falta = {m: 200 - len(v) for m, v in S['marcas'].items() if len(v) < 200}
ws = wb.create_sheet('Como foi feito')
for t in [
 'Como esta lista foi feita (9/10/2026)', '',
 'Onde buscamos',
 '• Google Maps, marca por marca (BMW, Mercedes, Audi, Porsche, Land Rover, Jaguar, Volvo e "importados/alemães"), em 32 regiões: capitais, cidades grandes do interior e bairros nobres e de oficinas de SP, Rio, BH, Porto Alegre, Curitiba e Brasília. Rolamos cada lista até o fim.',
 '• Ficaram de fora: concessionárias e autorizadas, redes e franquias (Bosch Car Service, Pit Stop, Midas, DPaschoal etc.), oficinas de moto, caminhão e diesel, lojas de peças, estética, lava-jato e quem tem o mesmo site em mais de 3 fichas.',
 '',
 'Como decidimos a marca de cada oficina',
 '• A marca aparece no nome, no site, na ficha do Google ou nas avaliações dos clientes (coluna "Por que é dessa marca"). Cada oficina aparece em uma marca só: a que ela mais cita. Oficinas de importados sem marca escrita só entram para completar uma marca, quando estão no topo da busca dela.',
 '',
 'Site',
 '• Sem site (10), só Instagram/Facebook/WhatsApp (9), site fora do ar (9), ou site com 3 ou mais problemas: ano velho no rodapé, não adaptado ao celular, lento, tecnologia antiga, sem HTTPS, sem prévia no WhatsApp, sem WhatsApp/formulário, links quebrados, erros ao abrir no celular, página mais larga que a tela.',
 '',
 'Sinais para vender o sistema',
 '• Sem agendamento on-line (nem botão de agendar no Google), site sem pedido de orçamento, atendimento só por WhatsApp/telefone, avaliações recentes que falam de orçamento, demora/prazo, retorno/contato ou preço, e volume de avaliações nos últimos 3 meses.',
 '',
 'Filtros: avaliação no Google nos últimos 12 meses, nota 3,8 ou mais, telefone. Telefones só para ligação comercial profissional.',
 f'Total: {tot} oficinas.' + (' Faltaram: ' + ', '.join(f'{m} {n}' for m, n in falta.items()) + ' (não existem mais oficinas independentes ativas com site ruim nessas marcas nas cidades buscadas).' if falta else ''),
 'Reservas: oficinas que sobraram (marca já completa ou evidência mais fraca).']:
    ws.append([t])
ws.column_dimensions['A'].width = 150
for c in ws['A']: c.alignment = Alignment(wrap_text=True)
ws['A1'].font = Font(bold=True, size=14)
for c in ws['A']:
    if c.value in ('Onde buscamos', 'Como decidimos a marca de cada oficina', 'Site', 'Sinais para vender o sistema'): c.font = Font(bold=True)
wb.save(xlsx)
os.makedirs(docsdir, exist_ok=True)
docs = []
for m, v in S['marcas'].items():
    k = CHAVE[m]
    for i, r in enumerate(v, 1):
        docs.append({'id': f'{k}-{i:03d}', 'data': {
            'lista': 'premium-' + k, 'nome': r['nome'], 'site': r['site'] if r['situacao'] != 'sem site' else '', 'cidade': r['cidade'],
            'segmento': f'Oficina {m}', 'status': 'novo', 'rank': i, 'triagem': '',
            'notas': f'{r["situacao"].capitalize()}. ' + ('Problemas: ' + '; '.join(r['P']) + '. ' if r['P'] else '') + 'Para o sistema: ' + '; '.join(r['S']) + '.',
            'contato': {'telefone': tel(r), 'email': r.get('email') or '', 'whatsapp': wa(r)},
            'extra': {'notaGoogle': r['nota_google'], 'avaliacoes': r['avaliacoes'], 'notaSite': r['nota'], 'maps': r['maps'],
                      'endereco': r.get('endereco') or '', 'bairro': ', '.join(r['bairros'] or []), 'ultimaAvaliacao': r['ult_txt'],
                      'problemas': r['P'], 'marca': m, 'porQueMarca': r['evidencia'], 'situacaoSite': r['situacao'],
                      'sinaisSistema': r['S'], 'trechosAvaliacoes': r['trechos']},
            'criadoEm': 1791600000000 - len(docs)}})
for j in range(0, len(docs), 50):
    json.dump([{'op': 'set', 'collection': 'empresas', 'doc_id': d['id'], 'data': d['data']} for d in docs[j:j + 50]],
              open(f'{docsdir}/lote-{j // 50:02d}.json', 'w'), ensure_ascii=False)
print('planilha ok,', len(docs), 'documentos')
