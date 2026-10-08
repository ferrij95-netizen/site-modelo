"""Gera a planilha e os documentos do canvas a partir da seleção (final.py)."""
import json, sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
S = json.load(open(sys.argv[1])); xlsx = sys.argv[2]; docsdir = sys.argv[3]
COLS = [('#', 5), ('Imobiliária', 34), ('Cidade', 18), ('Bairro', 22), ('Site', 34), ('Nota de site ruim (0-10)', 12),
        ('Problemas encontrados no site', 70), ('Última avaliação no Google', 16), ('Avaliações', 10), ('Nota Google', 9),
        ('Telefone', 17), ('WhatsApp', 15), ('E-mail', 28), ('Endereço', 45), ('Google Maps', 30)]
def wa(r):
    w = r.get('whatsapp') or ''
    return ('+' + w) if w else ''
def linha(i, r):
    return [i, r['nome'], r['cidade'], ', '.join(r['bairros']), r['site'], r['nota'], '; '.join(r['P']), r['ult_txt'],
            r['avaliacoes'], r['nota'] and r.get('nota_google'), r.get('telefone') or r.get('tel_site', ''), wa(r), r.get('email', ''),
            r.get('endereco', ''), r['maps']]
wb = Workbook()
for nome, lst in (('100 imobiliárias', S['top']), ('Reservas', S['res'])):
    ws = wb.active if nome.startswith('100') else wb.create_sheet(nome)
    ws.title = nome
    ws.append([c for c, _ in COLS])
    for i, r in enumerate(lst, 1):
        l = linha(i, r); l[9] = r.get('nota_g')
        ws.append(l)
        ws.cell(i + 1, 5).hyperlink = r['site']; ws.cell(i + 1, 15).hyperlink = r['maps']; ws.cell(i + 1, 15).value = 'abrir no Maps'
    for j, (_, w) in enumerate(COLS, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
        c = ws.cell(1, j); c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='182644')
        c.alignment = Alignment(wrap_text=True, vertical='center')
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'C2'; ws.auto_filter.ref = ws.dimensions
ws = wb.create_sheet('Como foi feito')
for t in [
 'Como esta lista foi feita (8/10/2026)',
 '',
 'Por que não é a lista que todo mundo tem',
 '• Quem roda raspador comum busca "imobiliária São Paulo" e pega os 20 primeiros. Nós rodamos essas buscas genéricas só para MARCAR esses nomes, e tiramos todos eles da lista (250 excluídos).',
 '• A busca de verdade foi bairro a bairro: 170 bairros de 14 cidades grandes (bairros residenciais e de classe média, não só os nobres), rolando a lista do Maps até o fim. Foram 15.942 fichas.',
 '• Tiramos redes e franquias (RE/MAX, Lopes, Century 21, Apolar, Auxiliadora, Foxter, QuintoAndar, Loft etc.), construtoras, portais e quem tem mais de 800 avaliações.',
 '',
 'Sinal de movimento recente',
 '• Para cada candidata abrimos a ficha do Google e lemos as avaliações. Só entra quem recebeu avaliação nos últimos 6 meses (coluna "Última avaliação").',
 '',
 'Como o site foi checado',
 '• Robô simples (10 pontos): rodapé com ano velho, não adaptado para celular, lento, tecnologia antiga (jQuery 1.x, Flash, frames), sem HTTPS, sem prévia ao compartilhar no WhatsApp, sem WhatsApp/formulário, conteúdo parado, links quebrados, sem busca de imóveis.',
 '• Navegador de verdade no tamanho de celular: erros de programação na página, imagens e arquivos que não carregam, página mais larga que a tela, demora para abrir, aviso de "site não seguro".',
 '• A nota é o número de problemas encontrados (0 a 10). Entrou quem tem 4 ou mais. No máximo 10 por cidade, para espalhar.',
 '',
 'Reservas: as que sobraram com nota 4+ e as ativas com nota 3.',
 '',
 'Limites: Instagram, OLX/ZAP e CRECI bloqueiam robôs, então o sinal de movimento é a avaliação no Google. Telefones vêm da ficha do Google ou do próprio site.']:
    ws.append([t])
ws.column_dimensions['A'].width = 140
for c in ws['A']: c.alignment = Alignment(wrap_text=True)
ws['A1'].font = Font(bold=True, size=14)
for k in (3, 8, 11): ws.cell(k, 1).font = Font(bold=True)
wb.save(xlsx)
# documentos do canvas (lista imobiliarias)
os.makedirs(docsdir, exist_ok=True)
docs = []
for i, r in enumerate(S['top'], 1):
    docs.append({'id': f'imob-{i:03d}', 'data': {
        'lista': 'imobiliarias', 'nome': r['nome'], 'site': r['site'], 'cidade': r['cidade'], 'segmento': r.get('categoria') or 'Imobiliária',
        'notas': 'Problemas no site: ' + '; '.join(r['P']), 'status': 'novo', 'rank': i, 'triagem': '',
        'contato': {'telefone': r.get('telefone') or r.get('tel_site', ''), 'email': r.get('email', ''), 'whatsapp': wa(r)},
        'extra': {'notaGoogle': r.get('nota_g'), 'avaliacoes': r['avaliacoes'], 'notaSite': r['nota'], 'maps': r['maps'],
                  'endereco': r.get('endereco', ''), 'bairro': ', '.join(r['bairros']), 'ultimaAvaliacao': r['ult_txt'], 'problemas': r['P']},
        'criadoEm': 1791500000000 - i}})
for k in range(0, len(docs), 50):
    json.dump([{'op': 'set', 'collection': 'empresas', 'doc_id': d['id'], 'data': d['data']} for d in docs[k:k + 50]],
              open(f'{docsdir}/lote-{k // 50}.json', 'w'), ensure_ascii=False)
print('ok', len(docs))
